"""
repair_gt_excel_dates.py — undo raw Excel date cells in ground-truth xlsx files, checked
against each paper's own text (the companion of repair_gt_dates.py, which undoes dates that
were already turned back into 'M.DD' numbers).

How the damage looks: a value typed as '1.2' (or '1-2', '1/2') was auto-converted by Excel into
a date cell (2 January of the year it was typed); openpyxl reads it as a datetime and the scorer
compares '2025-01-02 00:00:00' with the paper's '1.2'. The cell alone cannot tell whether the
typist meant month.day or day.month, or which separator; the paper can.

A measurementValue cell is repaired only when ALL hold:
  * it is a date cell (a datetime at midnight);
  * its trait is copied verbatim from the paper: >= 80% of the trait's other (undamaged) values
    occur in their own species' table row (repair_gt_dates._verbatim_traits' rule);
  * exactly one candidate original occurs in THIS species' own table row, trying 'M.D', 'D.M',
    'M-D', 'D-M', 'M/D', 'D/M' (and the zero-padded day/month forms). Several -> left alone.
A value never typed by hand (a real date, e.g. a sampling date) has no candidate in the row and
is left alone.

    python repair_gt_excel_dates.py <papers root> [paper ...]            # dry run: the log only
    python repair_gt_excel_dates.py <papers root> [paper ...] --apply    # also rewrites the xlsx

Before rewriting, the original xlsx is copied to data/backup/<root name>/<paper>/ (never
overwritten). Every change, and every date cell left alone with the reason, is logged in
data/<root name>_gt_excel_date_repairs.csv.
"""

from __future__ import annotations

import csv
import datetime as dt
import shutil
import sys
from pathlib import Path

import openpyxl

from mineru_extract import classify_folder
from repair_gt_dates import _in_text, _species_windows

ROOT = Path(__file__).resolve().parent.parent


def _is_date(v) -> bool:
    return isinstance(v, dt.datetime) and (v.hour, v.minute, v.second) == (0, 0, 0)


def _candidates(d: dt.datetime) -> list:
    m, day = d.month, d.day
    out = []
    for a, b in ((m, day), (day, m)):
        for sep in (".", "-", "/"):
            for x in (f"{a}{sep}{b}", f"{a}{sep}{b:02d}", f"{a:02d}{sep}{b:02d}"):
                if x not in out:
                    out.append(x)
    return out


def _value_str(v) -> str:
    return f"{v:.2f}" if isinstance(v, float) else str(v).strip()


def _verbatim_traits(wb, text, min_share=0.8, min_n=3):
    """{(sheet, measurementType)} whose undamaged values occur in their own species' row."""
    ok = set()
    for ws in wb.worksheets:
        hdr = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
        vcols = [i for i, h in enumerate(hdr) if h.startswith("measurementValue")]
        if not vcols or "measurementType" not in hdr or "verbatimIdentification" not in hdr:
            continue
        ti, si = hdr.index("measurementType"), hdr.index("verbatimIdentification")
        seen, found, windows = {}, {}, {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            trait, sp = row[ti], row[si]
            for vi in vcols:
                val = row[vi] if vi < len(row) else None
                if val is None or trait is None or _is_date(val):
                    continue
                v = _value_str(val)
                if not v:
                    continue
                if sp not in windows:
                    windows[sp] = _species_windows(sp, text)
                seen[trait] = seen.get(trait, 0) + 1
                found[trait] = found.get(trait, 0) + bool(windows[sp] and _in_text(v, windows[sp]))
        for trait, n in seen.items():
            if n >= min_n and found.get(trait, 0) / n >= min_share:
                ok.add((ws.title, trait))
    return ok


def repair_paper(folder: Path, apply: bool, backup: Path):
    _, gt, _ = classify_folder(folder)
    full = folder / f"{folder.name}_full.md"
    if not gt or not full.exists():
        return []
    text = full.read_text(encoding="utf-8", errors="replace")
    wb = openpyxl.load_workbook(gt)
    trusted = _verbatim_traits(wb, text)
    log, windows, changed = [], {}, False
    for ws in wb.worksheets:
        hdr = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
        vcols = [i for i, h in enumerate(hdr) if h.startswith("measurementValue")]
        if not vcols:
            continue
        ti = hdr.index("measurementType") if "measurementType" in hdr else None
        si = hdr.index("verbatimIdentification") if "verbatimIdentification" in hdr else None
        for row in ws.iter_rows(min_row=2):
            for vi in vcols:
                cell = row[vi]
                if not _is_date(cell.value):
                    continue
                sp = row[si].value if si is not None else ""
                trait = row[ti].value if ti is not None else ""
                rec = {"paper": folder.name, "file": Path(gt).name, "cell": f"{ws.title}!{cell.coordinate}",
                       "measurementType": trait, "species": sp, "old": cell.value.strftime("%Y-%m-%d"),
                       "new": "", "status": ""}
                if (ws.title, trait) not in trusted:
                    rec["status"] = "left: trait not copied verbatim from the paper"
                    log.append(rec)
                    continue
                if sp not in windows:
                    windows[sp] = _species_windows(sp, text)
                local = windows[sp]
                found = [c for c in _candidates(cell.value) if local and _in_text(c, local)]
                if len(found) != 1:
                    rec["status"] = ("left: no candidate in the species' row" if not found
                                     else f"left: several candidates {found}")
                    log.append(rec)
                    continue
                rec["new"], rec["status"] = found[0], "repaired"
                log.append(rec)
                if apply:
                    cell.value = found[0]
                    cell.number_format = "@"
                    changed = True
    if changed:
        dest = backup / folder.name / Path(gt).name
        dest.parent.mkdir(parents=True, exist_ok=True)
        if not dest.exists():
            shutil.copy2(gt, dest)
        wb.save(gt)
    return log


def main(argv):
    apply = "--apply" in argv
    args = [a for a in argv if not a.startswith("--")]
    root = Path(args[0]).resolve()
    names = args[1:]
    folders = [root / n for n in names] if names else sorted(p for p in root.iterdir() if p.is_dir())
    backup = ROOT / "data" / "backup" / root.name
    allr = []
    for f in folders:
        try:
            r = repair_paper(f, apply, backup)
        except Exception as e:
            print(f"skip {f.name}: {e}")
            continue
        if r:
            n = sum(x["status"] == "repaired" for x in r)
            print(f"{f.name:<24} {n:>5} repaired, {len(r) - n:>5} date cells left")
            allr += r
    log = ROOT / "data" / f"{root.name}_gt_excel_date_repairs.csv"
    with open(log, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["paper", "file", "cell", "measurementType", "species",
                                           "old", "new", "status"])
        w.writeheader()
        w.writerows(allr)
    n = sum(x["status"] == "repaired" for x in allr)
    print(f"\n{n} value(s) {'repaired' if apply else 'to repair (dry run)'}, "
          f"{len(allr) - n} date cells left, in {len({x['paper'] for x in allr})} paper(s) -> {log}")


if __name__ == "__main__":
    main(sys.argv[1:])
