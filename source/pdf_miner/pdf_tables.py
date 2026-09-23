"""
pdf_tables.py — extract ONLY real tables from a PDF, MinerU-style.

Instead of trusting a converter's markdown (which pours figures into <table> tags),
we decide what is a table VISUALLY, up front, then read only those regions:

    1. render each PDF page to an image
    2. DocLayout-YOLO on the image -> boxes labelled table / figure / text / ...
    3. keep only the `table` boxes (figures are dropped at the source)
    4. crop each table box out of the PDF (pdfplumber page.crop, in PDF points)
    5. cluster the cropped words into a grid
    6. write one sheet per table into data/structured/<Paper>.xlsx

Cross-page tables: YOLO is per-page, so a table split across a page break is
detected as two boxes. A light stitch merges a box touching the bottom margin with
the next page's box touching the top margin when their column counts match.

Usage:
    python pdf_tables.py --paper Barber2017
    python pdf_tables.py                    # all papers -> data/structured/
    python pdf_tables.py --resume
"""
from __future__ import annotations

import argparse
import glob
import os
import sys

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DATA_ROOT = os.path.join(REPO_ROOT, "data", "un_processed_papers")
STRUCT_ROOT = os.path.join(REPO_ROOT, "data", "structured")

RENDER_SCALE = 2.0          # page -> image zoom (bbox is divided by this to get PDF pts)
TABLE_CONF = 0.35           # min YOLO confidence for a `table` box
_MODEL = None


def _weights_path() -> str:
    from huggingface_hub import hf_hub_download
    return hf_hub_download(repo_id="juliozhao/DocLayout-YOLO-DocStructBench",
                           filename="doclayout_yolo_docstructbench_imgsz1024.pt")


def _model():
    global _MODEL
    if _MODEL is None:
        from doclayout_yolo import YOLOv10
        _MODEL = YOLOv10(_weights_path())
    return _MODEL


# ---------------------------------------------------------------------------
# 1-3. detect table boxes (in PDF points) per page
# ---------------------------------------------------------------------------

def detect_table_boxes(pdf_path: str):
    """Return [(page_no, (x0, top, x1, bottom) in PDF points, conf), ...] for every
    region DocLayout-YOLO labels `table`."""
    import fitz
    import numpy as np
    model = _model()
    boxes = []
    doc = fitz.open(pdf_path)
    for pno in range(len(doc)):
        pix = doc[pno].get_pixmap(matrix=fitz.Matrix(RENDER_SCALE, RENDER_SCALE))
        img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)[:, :, :3]
        res = model.predict(img, imgsz=1024, conf=0.2, device="cuda", verbose=False)[0]
        for b, cls, cf in zip(res.boxes.xyxy.tolist(), res.boxes.cls.tolist(), res.boxes.conf.tolist()):
            if res.names[int(cls)] == "table" and cf >= TABLE_CONF:
                x0, y0, x1, y1 = (v / RENDER_SCALE for v in b)
                boxes.append((pno, (x0, y0, x1, y1), round(cf, 2)))
    doc.close()
    return boxes


# ---------------------------------------------------------------------------
# 4-5. crop a box and cluster its words into a grid
# ---------------------------------------------------------------------------

def _rows_of(words, y_tol):
    rows, cur, last = [], [], None
    for w in sorted(words, key=lambda w: (w["top"], w["x0"])):
        if last is None or abs(w["top"] - last) <= y_tol:
            cur.append(w)
        else:
            rows.append(cur)
            cur = [w]
        last = w["top"]
    if cur:
        rows.append(cur)
    return [sorted(r, key=lambda w: w["x0"]) for r in rows]


def _column_edges(rows):
    """Column separators = x0 positions recurring across rows (same idea as the
    geometry backend). Returns sorted x-midpoints between column starts."""
    from collections import Counter
    bins, n = Counter(), 0
    for r in rows:
        n += 1
        for b in {round(w["x0"] / 3) * 3 for w in r}:
            bins[b] += 1
    peaks = sorted(b for b, c in bins.items() if c >= max(2, n * 0.3))
    edges = []
    for p in peaks:
        if not edges or p - edges[-1] >= 8:
            edges.append(p)
    return [(edges[i] + edges[i + 1]) / 2 for i in range(len(edges) - 1)]


def _assign(row, edges):
    cells = [""] * (len(edges) + 1)
    for w in row:
        col = next((i for i, e in enumerate(edges) if w["x0"] < e), len(edges))
        cells[col] = (cells[col] + " " + w["text"]).strip()
    return cells


def crop_to_grid(page, bbox):
    """pdfplumber page cropped to `bbox` (PDF points) -> aligned grid of strings."""
    x0, top, x1, bottom = bbox
    x0 = max(0, x0); top = max(0, top)
    x1 = min(page.width, x1); bottom = min(page.height, bottom)
    if x1 - x0 < 5 or bottom - top < 5:
        return []
    region = page.crop((x0, top, x1, bottom))
    words = region.extract_words(use_text_flow=False, keep_blank_chars=False)
    if not words:
        return []
    heights = sorted(w["bottom"] - w["top"] for w in words)
    y_tol = max(3.0, heights[len(heights) // 2] * 0.6)
    rows = _rows_of(words, y_tol)
    edges = _column_edges(rows)
    grid = [_assign(r, edges) for r in rows]
    width = max((len(r) for r in grid), default=0)
    return [r + [""] * (width - len(r)) for r in grid if any(c.strip() for c in r)]


# ---------------------------------------------------------------------------
# 6. one paper -> one workbook (index + a sheet per table)
# ---------------------------------------------------------------------------

def _stitch(tables):
    """Merge a table whose box touches the page bottom with the next page's box
    touching the top, when column counts match (a page-continued table)."""
    out = []
    for t in tables:
        if out:
            prev = out[-1]
            prev_bottom = prev["bbox"][3] >= prev["page_h"] * 0.9
            cur_top = t["bbox"][1] <= t["page_h"] * 0.12
            same_cols = prev["grid"] and t["grid"] and len(prev["grid"][0]) == len(t["grid"][0])
            if prev["pdf"] == t["pdf"] and t["page"] == prev["page"] + 1 and prev_bottom and cur_top and same_cols:
                prev["grid"] += t["grid"]
                prev["stitched"] = prev.get("stitched", 1) + 1
                continue
        out.append(t)
    return out


def extract_paper(pdf_path: str):
    """All table grids from one PDF, in reading order, with stitching."""
    import pdfplumber
    boxes = detect_table_boxes(pdf_path)
    tables = []
    with pdfplumber.open(pdf_path) as pdf:
        for pno, bbox, conf in sorted(boxes, key=lambda x: (x[0], x[1][1])):
            page = pdf.pages[pno]
            grid = crop_to_grid(page, bbox)
            if grid and len(grid) >= 2:
                tables.append({"pdf": os.path.basename(pdf_path), "page": pno,
                               "bbox": bbox, "conf": conf, "grid": grid,
                               "page_h": page.height})
    return _stitch(tables)


def write_workbook(name: str, tables: list, out_path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side
    thin = Side(style="thin", color="C0C0C0")
    border = Border(thin, thin, thin, thin)
    wb = Workbook()
    idx = wb.active
    idx.title = "index"
    idx.append(["sheet", "source_pdf", "page", "conf", "rows", "cols", "stitched_pages"])
    for c in idx[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for i, t in enumerate(tables):
        sheet = f"t{i:02d}_p{t['page']}"[:31]
        idx.append([sheet, t["pdf"], t["page"], t["conf"],
                    len(t["grid"]), len(t["grid"][0]), t.get("stitched", 1)])
        ws = wb.create_sheet(title=sheet)
        for r, row in enumerate(t["grid"], 1):
            for c, val in enumerate(row, 1):
                cell = ws.cell(r, c, val or None)
                cell.border = border
                if r == 1:
                    cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
    if len(wb.sheetnames) == 1:               # no tables found: leave a note
        idx.append(["(no tables detected)"])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)


def paper_pdfs(paper_dir: str):
    """Main PDF first, then supplements (so table order is sensible)."""
    name = os.path.basename(paper_dir)
    pdfs = sorted(glob.glob(os.path.join(paper_dir, "*.pdf")))
    pdfs.sort(key=lambda p: (os.path.basename(p) != f"{name}.pdf",))
    return pdfs


def main(argv=None):
    global TABLE_CONF
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--paper", help="paper folder name (default: all)")
    ap.add_argument("--resume", action="store_true", help="skip papers whose xlsx exists")
    ap.add_argument("--limit", type=int)
    ap.add_argument("--conf", type=float, default=TABLE_CONF, help="min table confidence")
    args = ap.parse_args(argv)
    TABLE_CONF = args.conf

    papers = ([os.path.join(DATA_ROOT, args.paper)] if args.paper
              else sorted(d for d in glob.glob(os.path.join(DATA_ROOT, "*")) if os.path.isdir(d)))
    if args.limit:
        papers = papers[:args.limit]

    import time as _t
    for pn, pdir in enumerate(papers, 1):
        name = os.path.basename(pdir)
        out = os.path.join(STRUCT_ROOT, f"{name}.xlsx")
        if args.resume and os.path.exists(out):
            print(f"[{pn}/{len(papers)}] {name}: skip", flush=True)
            continue
        t0 = _t.time()
        tables = []
        try:
            for pdf in paper_pdfs(pdir):
                tables += extract_paper(pdf)
        except Exception as e:
            print(f"[{pn}/{len(papers)}] {name}: ERROR {e}", flush=True)
            continue
        write_workbook(name, tables, out)
        print(f"[{pn}/{len(papers)}] {name}: {len(tables)} tables ({_t.time()-t0:.0f}s)", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
