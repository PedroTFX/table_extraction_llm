"""Correct Noriega2023's GT values to the values the markdown prints (user request,
2026-10-19). Source: Noriega2023_S1.xlsx.md, tables S6 (TRAITS.raw, one row per specimen
and trait) and S5 (TRAITS.summary, one row per species).

  individual rows  per (species, trait): the GT list must hold the S6 values of that
                   species and trait code (codes matched ignoring case: 'mTL' = 'MTL');
                   GT values the source list holds stay, the others are replaced in order
                   by the source values the GT lacks (same length only; else reported)
  mean rows        the S5 value of that species and trait column
  mode rows        (relocation size, decoded from S5 codes by the annotator) left as they are
Groups whose other columns differ between rows are left. Every changed cell is appended to
output/gt_value_fixes.csv; the GT before the change is saved in data/backup/Noriega2023/.

    python fix_gt_noriega.py [--dry]
"""
import collections
import csv
import html
import re
import shutil
import sys

import openpyxl

import _paths

sys.path.insert(0, str(_paths.REPO / "source"))
import evaluate as ev  # noqa: E402

PAPER = "Noriega2023"
GT = _paths.PAPERS / PAPER / f"{PAPER}.xlsx"
MD = _paths.PAPERS / PAPER / f"{PAPER}_S1.xlsx.md"
BACKUP = _paths.REPO / "data" / "backup" / PAPER / f"{PAPER}_2026-10-19.xlsx"
LOG = _paths.REPO / "output" / "gt_value_fixes.csv"


def sp(s):
    return " ".join(str(s or "").replace("_", " ").split())


def norm(v):
    return " ".join(str(v if v is not None else "").split())


def md_tables(text):
    """[(header, rows)] of the <table>s in the markdown, cells as printed."""
    out = []
    for m in re.finditer(r"<table>(.*?)</table>", text, re.S):
        rows = []
        for tr in re.findall(r"<tr>(.*?)</tr>", m.group(1), re.S):
            rows.append([html.unescape(re.sub(r"<[^>]+>", "", c)).strip()
                         for c in re.findall(r"<td[^>]*>(.*?)</td>", tr, re.S)])
        if rows:
            out.append((rows[0], rows[1:]))
    return out


def find(tables, first_cols):
    for hdr, rows in tables:
        if hdr[:len(first_cols)] == first_cols:
            return hdr, rows
    raise SystemExit(f"table {first_cols} not found in the markdown")


def value_map(gt_lists, src_lists):
    """GT trait name -> source trait codes (lower case), by most values in common."""
    out = {}
    for ty in {t for _s, t in gt_lists}:
        score = collections.Counter()
        for (s, t), vals in gt_lists.items():
            if t != ty:
                continue
            for (s2, code), v2 in src_lists.items():
                if s2 == s:
                    score[code] += len(set(vals) & set(v2))
        if score:
            best = score.most_common(1)[0][0]
            out[ty] = best
    return out


def main(dry):
    tables = md_tables(MD.read_text(encoding="utf-8"))
    h6, r6 = find(tables, ["Site Code", "Dung relocation strategy", "Subfamily", "Species"])
    i_sp, i_tr, i_me = h6.index("Species"), h6.index("Trait"), h6.index("Measurement")
    s6 = collections.defaultdict(list)                   # (species, code lower) -> [printed]
    for r in r6:
        if len(r) > i_me and r[i_sp]:
            s6[(sp(r[i_sp]), r[i_tr].strip().lower())].append(r[i_me])
    h5, r5 = find(tables, ["Trait", "Biomass", "EL"])
    s5 = {}
    for r in r5:
        if r and r[0] and r[0] != "Trait Type":
            for h, v in zip(h5[1:], r[1:]):
                s5[(sp(r[0]), h.strip().lower())] = v

    wb = openpyxl.load_workbook(GT)
    ws = wb.active
    hdr = [str(c.value or "").strip() for c in ws[1]]
    c_sp, c_ty = hdr.index("verbatimIdentification"), hdr.index("measurementType")
    c_va, c_st = hdr.index("measurementValue"), hdr.index("measurementStatistic")
    groups = collections.defaultdict(list)                # (species, type, stat) -> [row numbers]
    others = collections.defaultdict(set)
    for row in ws.iter_rows(min_row=2):
        vals = [c.value for c in row]
        if not vals[c_sp]:
            continue
        k = (sp(vals[c_sp]), norm(vals[c_ty]), norm(vals[c_st]))
        groups[k].append(row[0].row)
        others[k].add(tuple(norm(v) for i, v in enumerate(vals) if i != c_va))

    def gt_vals(k):
        return [norm(ws.cell(row=n, column=c_va + 1).value) for n in groups[k]]

    ind = {(s, t): gt_vals((s, t, st)) for (s, t, st) in groups if st == "individual"}
    tmap = value_map(ind, s6)
    mean = {(s, t): gt_vals((s, t, st)) for (s, t, st) in groups if st == "mean"}
    mmap = value_map(mean, {(s, h): [v] for (s, h), v in s5.items()})
    print("individual: GT trait -> S6 code:", tmap)
    print("mean:       GT trait -> S5 column:", mmap)

    changes, skipped = [], collections.Counter()
    col = openpyxl.utils.get_column_letter(c_va + 1)
    for (s, t, st), rows in groups.items():
        if len(others[(s, t, st)]) > 1:
            skipped["other columns differ"] += len(rows)
            continue
        if st == "individual":
            code = tmap.get(t)
            src = s6.get((s, code)) if code else None
            if not src or len(src) != len(rows):
                skipped[f"count differs or no source"] += len(rows)
                continue
            new_vals = src
            why = f"S6 printed value (species {s}, trait {code}, specimen order)"
        elif st == "mean":
            h = mmap.get(t)
            if not h or (s, h) not in s5 or len(rows) != 1:
                skipped["no S5 value"] += len(rows)
                continue
            new_vals = [s5[(s, h)]]
            why = f"S5 printed value (species {s}, column {h})"
        else:
            skipped[f"{st} rows left"] += len(rows)
            continue
        # minimal change (the scorer compares rows regardless of order): a GT value that the
        # source list holds stays; the others take, in order, the source values the GT lacks
        # values compared as the scorer compares them ('1' = '1.0'): no cosmetic changes
        olds = [norm(ws.cell(row=n, column=c_va + 1).value) for n in rows]
        have = collections.Counter(ev.normalize_value(v) for v in new_vals)
        keep = []
        for old in olds:
            k = ev.normalize_value(old)
            keep.append(have[k] > 0)
            if have[k] > 0:
                have[k] -= 1
        missing = []
        for v in new_vals:                       # source values the GT lacks, in source order
            k = ev.normalize_value(v)
            if have[k] > 0:
                have[k] -= 1
                missing.append(v)
        for n, old, kept in zip(rows, olds, keep):
            if kept:
                continue
            new = missing.pop(0)
            changes.append((f"{col}{n}", old, new, why + "; the list now equals the source list "
                            "(pairs by position, not one-to-one corrections)"))
            ws.cell(row=n, column=c_va + 1).value = new
    print(f"cells to change: {len(changes)}; left: {dict(skipped)}")
    kinds = collections.Counter("mean (S5)" if "S5" in w else "individual (S6)" for _c, _o, _n, w in changes)
    print("   by row type:", dict(kinds))
    for c in changes[:8]:
        print("   ", c)
    if dry:
        return
    BACKUP.parent.mkdir(parents=True, exist_ok=True)
    if not BACKUP.exists():
        shutil.copy(GT, BACKUP)
    wb.save(GT)
    with open(LOG, "a", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        for cell, old, new, why in changes:
            w.writerow([PAPER, cell, old, new, f"GT value -> markdown value; {why}"])
    print(f"saved {GT.name}; backup {BACKUP}; {len(changes)} cells logged in {LOG.name}")


if __name__ == "__main__":
    main("--dry" in sys.argv)
