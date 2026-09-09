from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook

# Parent folder containing all paper folders
parent_folder = Path(r"C:/Users/Pedro Trindade/Documents/github/table_extraction_llm/data/un_processed_papers - Copy")   # <-- change this

total_files = 0
total_repairs = 0

for folder in parent_folder.iterdir():

    if not folder.is_dir():
        continue

    # Process every xlsx in this folder
    for excel_file in folder.glob("*.xlsx"):

        print(f"\nProcessing {excel_file.name}")

        wb = load_workbook(excel_file)

        modified = False
        repaired_here = 0

        for ws in wb.worksheets:

            for row in ws.iter_rows():

                for cell in row:

                    if isinstance(cell.value, datetime):

                        month = cell.value.month
                        day = cell.value.day

                        repaired = float(f"{month}.{day:02d}")

                        print(
                            f"  {ws.title} {cell.coordinate}: "
                            f"{cell.value.strftime('%d-%b')} -> {repaired:.2f}"
                        )

                        # Replace date with numeric value
                        cell.style = "Normal"
                        cell.value = repaired
                        cell.number_format = "0.00"

                        modified = True
                        repaired_here += 1

        if modified:
            wb.save(excel_file)
            print(f"  Saved ({repaired_here} repairs)")
            total_repairs += repaired_here

        total_files += 1

print("\n--------------------------------")
print(f"Files processed : {total_files}")
print(f"Cells repaired  : {total_repairs}")
print("--------------------------------")