from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook

# Parent folder containing all paper folders
parent_folder = Path(r"C:/Users/Pedro Trindade/Documents/github/table_extraction_llm/data/un_processed_papers")   # <-- change this

for folder in parent_folder.iterdir():

    if not folder.is_dir():
        continue

    excel_file = folder / f"{folder.name}.xlsx"

    if not excel_file.exists():
        continue

    print(f"\nProcessing {excel_file.name}")

    wb = load_workbook(excel_file)

    modified = False

    for ws in wb.worksheets:

        # Find the measurementValue column
        measurement_col = None

        for cell in ws[1]:
            if cell.value is None:
                continue

            header = str(cell.value).strip().replace("\xa0", "")

            if header == "measurementValue":
                measurement_col = cell.column
                break

        if measurement_col is None:
            continue

        for row in range(2, ws.max_row + 1):

            cell = ws.cell(row=row, column=measurement_col)

            if isinstance(cell.value, datetime):

                month = cell.value.month
                day = cell.value.day

                repaired = float(f"{month}.{day:02d}")

                print(
                    f"{ws.title} {cell.coordinate}: "
                    f"{cell.value.strftime('%d-%b')} -> {repaired}"
                )

                cell.value = repaired
                cell.number_format = "0.00"

                modified = True

    if modified:
        wb.save(excel_file)
        print("Saved.")
    else:
        print("No repairs needed.")